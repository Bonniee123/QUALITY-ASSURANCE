"""Dashboard views with analytics and summary data."""
import calendar
import csv
import json
from datetime import date, datetime, time, timedelta

from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.db.models.functions import TruncDate, TruncMonth, TruncWeek
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone

from accounts.decorators import repository_access_required
from accounts.permissions import is_admin, scope_documents_for_user, faculty_area_scope
from ai_processing.dashboard_utils import cluster_label_map
from documents.models import ActivityLog, Document
from documents.excel_export import (
    csv_safe_row, openpyxl_available, build_styled_sheet, excel_response,
)
from qa_mapping.models import QAProgram

DUPLICATE_REVIEW_STATUSES = ('possible', 'pending_check', 'confirmed_dup')


def _kpi_delta(current: int, previous: int) -> dict:
    """
    How a period compares with the one before it.

    An empty previous period used to return a flat "100.0" pointing up. That
    number was a placeholder, not a measurement: growth from zero is undefined,
    not a hundred per cent. On this archive every document was uploaded inside
    one twelve-minute window, so both the 30-day and the 7-day comparison had
    nothing behind them, and two cards sat side by side each claiming "up
    100.0%" -- a figure nobody could derive from anything on the page.

    That case now returns the direction 'new', and the card prints a word
    instead of a percentage. This applies to every role: the comparison is
    computed the same way whoever is looking.
    """
    if previous == 0:
        if current == 0:
            return {'pct': '0.0', 'dir': 'flat'}
        return {'pct': '', 'dir': 'new'}
    delta = ((current - previous) / previous) * 100.0
    direction = 'up' if delta > 0 else ('down' if delta < 0 else 'flat')
    return {'pct': f'{abs(delta):.1f}', 'dir': direction}


def _aware(d):
    if hasattr(d, 'tzinfo') and d.tzinfo is not None:
        return d
    if not isinstance(d, datetime):
        d = datetime.combine(d, time.min)
    return timezone.make_aware(d) if timezone.is_naive(d) else d


def _build_trend_buckets(period: str, now):
    today = timezone.localdate(now)
    if period == 'daily':
        starts = [today - timedelta(days=i) for i in range(29, -1, -1)]
        labels = [d.strftime('%b %d') for d in starts]
        return labels, starts, TruncDate('uploaded_at')
    if period == 'weekly':
        anchor = today - timedelta(days=today.weekday())
        starts = [anchor - timedelta(weeks=11 - i) for i in range(12)]
        labels = [d.strftime('%b %d') for d in starts]
        return labels, starts, TruncWeek('uploaded_at')
    starts = []
    cursor = today.replace(day=1)
    for _ in range(6):
        starts.append(cursor)
        prev_last = cursor - timedelta(days=1)
        cursor = prev_last.replace(day=1)
    starts.reverse()
    labels = [d.strftime('%b') for d in starts]
    return labels, starts, TruncMonth('uploaded_at')


def _bucket_counts(qs, trunc_expr, starts):
    rows = qs.annotate(b=trunc_expr).values('b').annotate(c=Count('id'))
    by_key = {}
    for row in rows:
        key = row['b']
        if hasattr(key, 'date'):
            key = key.date()
        by_key[key] = row['c']
    return [by_key.get(d, 0) for d in starts]


def _parse_calendar_month(request, now):
    cal_param = request.GET.get('cal', '').strip()
    today = timezone.localdate(now)
    if cal_param:
        try:
            parts = cal_param.split('-')
            year, month = int(parts[0]), int(parts[1])
            if 1 <= month <= 12 and 2000 <= year <= 2100:
                return year, month
        except (ValueError, IndexError, TypeError):
            pass
    return today.year, today.month


def _calendar_nav_query(request, year, month, period, current_program):
    if month == 1:
        prev_y, prev_m = year - 1, 12
    else:
        prev_y, prev_m = year, month - 1
    if month == 12:
        next_y, next_m = year + 1, 1  # was 12: "Next" from December went to December of the next year
    else:
        next_y, next_m = year, month + 1

    def _qs(y, m):
        params = [f'cal={y:04d}-{m:02d}', f'period={period}']
        if current_program:
            params.append(f'program={current_program.pk}')
        return '?' + '&'.join(params)

    return {
        'prev_url': _qs(prev_y, prev_m),
        'next_url': _qs(next_y, next_m),
        'prev_label': date(prev_y, prev_m, 1).strftime('%b %Y'),
        'next_label': date(next_y, next_m, 1).strftime('%b %Y'),
        'cal_param': f'{year:04d}-{month:02d}',
    }


def _build_calendar_weeks(year, month, upload_counts):
    cal = calendar.Calendar(firstweekday=0)
    weeks = []
    for week in cal.monthdatescalendar(year, month):
        row = []
        for day in week:
            count = upload_counts.get(day, 0)
            row.append({
                'date': day,
                'iso': day.isoformat(),
                'day': day.day,
                'in_month': day.month == month,
                'count': count,
                'has_uploads': count > 0,
            })
        weeks.append(row)
    return weeks


def _csv_export(now, total_documents, duplicate_review_count,
                monthly_labels, monthly_counts,
                file_format_labels, file_format_counts, program_overview,
                current_program):
    response = HttpResponse(content_type='text/csv')
    stamp = now.strftime('%Y%m%d_%H%M%S')
    suffix = f'_{current_program.code}' if current_program else ''
    response['Content-Disposition'] = f'attachment; filename="dashboard_summary{suffix}_{stamp}.csv"'

    w = csv.writer(response)
    w.writerow(['Section', 'Item', 'Value', 'Value 2', 'Value 3', 'Notes'])
    w.writerow(['Summary', 'Generated', now.strftime('%Y-%m-%d %H:%M:%S'), '', '', ''])
    if current_program:
        w.writerow(csv_safe_row(['Summary', 'Filtered by program', current_program.code, current_program.name, '', '']))

    w.writerow(['Key Metrics', 'Total documents', total_documents, '', '', ''])
    w.writerow(['Key Metrics', 'Duplicates to review', duplicate_review_count, '', '', ''])

    for label, count in zip(monthly_labels, monthly_counts):
        w.writerow(['Upload Trend', label, count, '', '', ''])

    for label, count in zip(file_format_labels, file_format_counts):
        w.writerow(['File Format Distribution', label, count, '', '', ''])

    for p in program_overview:
        w.writerow(csv_safe_row([
            'QA Programs',
            p['program'].code,
            p['program'].name,
            p['documents'],
            '',
            '',
        ]))
    return response


def _excel_export(now, metrics, monthly_labels, monthly_counts,
                  file_format_labels, file_format_counts, program_overview, current_program):
    """Build a styled, multi-sheet Excel workbook summarizing the dashboard."""
    from openpyxl import Workbook

    scope = current_program.code if current_program else 'All QA Documents'
    wb = Workbook()

    summary_rows = [
        ['Scope', scope],
        ['Total documents', metrics['total_documents']],
        ['Uploaded (last 7 days)', metrics['docs_last7']],
        ['Uploaded (last 30 days)', metrics['docs_last30']],
        ['Duplicates to review', metrics['duplicate_review_count']],
        ['Document clusters', metrics['cluster_count']],
        ['Analyzed (%)', metrics['ai_processed_pct']],
        ['Analyzed documents', metrics['ai_processed_count']],
    ]
    ws = wb.active
    ws.title = 'Summary'
    build_styled_sheet(ws, 'Dashboard Summary', ['Metric', 'Value'], summary_rows, generated_at=now)

    trend = wb.create_sheet('Upload Trend')
    build_styled_sheet(
        trend, 'Upload Trend', ['Period', 'Documents'],
        list(zip(monthly_labels, monthly_counts)), generated_at=now,
    )

    fmt = wb.create_sheet('File Formats')
    build_styled_sheet(
        fmt, 'File Format Distribution', ['Format', 'Documents'],
        list(zip(file_format_labels, file_format_counts)), generated_at=now,
    )

    progs = wb.create_sheet('QA Programs')
    build_styled_sheet(
        progs, 'QA Programs', ['Code', 'Program', 'Documents'],
        [[p['program'].code, p['program'].name, p['documents']] for p in program_overview],
        generated_at=now,
    )

    suffix = f'_{current_program.code}' if current_program else ''
    return excel_response(wb, f'dashboard_summary{suffix}_{now:%Y%m%d_%H%M%S}')


@login_required
@repository_access_required
def dashboard_home(request):
    """Main dashboard with summary cards, charts, calendar, and recent activity."""
    now = timezone.now()
    area_scope = faculty_area_scope(request.user)
    thirty_days_ago = now - timedelta(days=30)
    sixty_days_ago = now - timedelta(days=60)

    period = request.GET.get('period', 'monthly').lower()
    if period not in ('monthly', 'weekly', 'daily'):
        period = 'monthly'

    program_filter_id = request.GET.get('program', '').strip()
    current_program = None
    if program_filter_id.isdigit():
        current_program = QAProgram.objects.filter(pk=int(program_filter_id)).first()

    docs_qs = Document.objects.live().filter(is_archived=False)
    docs_qs = scope_documents_for_user(docs_qs, request.user)
    if current_program:
        docs_qs = docs_qs.filter(program=current_program)

    total_documents = docs_qs.count()
    duplicate_review_count = docs_qs.filter(
        duplicate_status__in=DUPLICATE_REVIEW_STATUSES,
    ).count()

    docs_last30 = docs_qs.filter(uploaded_at__gte=thirty_days_ago).count()
    docs_prev30 = docs_qs.filter(
        uploaded_at__gte=sixty_days_ago, uploaded_at__lt=thirty_days_ago,
    ).count()
    docs_delta = _kpi_delta(docs_last30, docs_prev30)

    # "This week" — last 7 days, with delta vs. the previous 7 days.
    seven_days_ago = now - timedelta(days=7)
    fourteen_days_ago = now - timedelta(days=14)
    docs_last7 = docs_qs.filter(uploaded_at__gte=seven_days_ago).count()
    docs_prev7 = docs_qs.filter(
        uploaded_at__gte=fourteen_days_ago, uploaded_at__lt=seven_days_ago,
    ).count()
    week_delta = _kpi_delta(docs_last7, docs_prev7)

    # Distinct AI cluster groups present in the current scope.
    cluster_count = (
        docs_qs.filter(cluster_label__isnull=False)
        .values('cluster_label')
        .distinct()
        .count()
    )

    # Analyzed = documents that have been through processing -- the Document
    # Analysis page's own definition, so the two pages give the same number.
    # This counted "has extracted text" under the old "AI-Processed" label.
    ai_processed_count = docs_qs.filter(is_processed=True).count()
    ai_processed_pct = round((ai_processed_count / total_documents) * 100) if total_documents else 0

    # 14-day daily series for the KPI sparklines.
    spark_days = [now.date() - timedelta(days=i) for i in range(13, -1, -1)]
    spark_counts = _bucket_counts(
        docs_qs.filter(uploaded_at__gte=_aware(spark_days[0])),
        TruncDate('uploaded_at'), spark_days,
    )

    trend_labels, trend_starts, trend_trunc = _build_trend_buckets(period, now)
    trend_window_start = _aware(trend_starts[0])

    monthly_counts = _bucket_counts(
        docs_qs.filter(uploaded_at__gte=trend_window_start),
        trend_trunc, trend_starts,
    )
    monthly_labels = trend_labels
    monthly_total_window = sum(monthly_counts)

    palette = [
        'hsl(180, 65%, 40%)', 'hsl(220, 70%, 55%)', 'hsl(265, 65%, 60%)',
        'hsl(330, 65%, 55%)', 'hsl(35, 90%, 55%)', 'hsl(150, 60%, 45%)',
    ]
    active_programs = list(QAProgram.objects.filter(is_active=True))
    program_lines = []
    for idx, prog in enumerate(active_programs):
        if current_program and prog.pk != current_program.pk:
            continue
        pcounts = _bucket_counts(
            docs_qs.filter(program=prog, uploaded_at__gte=trend_window_start),
            trend_trunc, trend_starts,
        )
        if sum(pcounts) == 0:
            continue
        program_lines.append({
            'label': prog.code or prog.name[:12],
            'data': pcounts,
            'color': prog.color or palette[idx % len(palette)],
        })

    if not program_lines and monthly_total_window > 0:
        program_lines.append({
            'label': 'All uploads' if not current_program else current_program.code,
            'data': monthly_counts,
            'color': 'hsl(180, 60%, 35%)',
        })

    recent_uploads = docs_qs.select_related('uploaded_by').order_by('-uploaded_at')[:10]
    # The feed shows archive events from anyone plus the viewer's own personal
    # events. Personal actions -- logins, page views, searches, messages -- are
    # not other people's business on a shared dashboard; they stay in the
    # administrator-only Audit Log, which keeps the complete record.
    from documents.activity_display import dashboard_feed

    activity_qs = ActivityLog.objects.select_related('user').order_by('-created_at')
    if area_scope is not None:
        activity_qs = activity_qs.filter(user=request.user)
    recent_activity = dashboard_feed(activity_qs, request.user, limit=15)

    file_format_data = (
        docs_qs.values('file_type')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    file_format_labels = [(row['file_type'] or 'OTHER').upper() for row in file_format_data]
    file_format_counts = [row['count'] for row in file_format_data]
    file_format_total = sum(file_format_counts)

    program_overview = []
    for prog in active_programs:
        if current_program and prog.pk != current_program.pk:
            continue
        # Counted from the same documents as the KPIs -- in scope and not archived
        # (the program's own count included archived versions).
        prog_count = docs_qs.filter(program=prog).count()
        program_overview.append({
            'program': prog,
            'documents': prog_count,
        })

    cal_year, cal_month = _parse_calendar_month(request, now)
    month_start = date(cal_year, cal_month, 1)
    if cal_month == 12:
        month_end = date(cal_year + 1, 1, 1) - timedelta(days=1)
    else:
        month_end = date(cal_year, cal_month + 1, 1) - timedelta(days=1)

    upload_rows = (
        docs_qs.filter(uploaded_at__date__gte=month_start, uploaded_at__date__lte=month_end)
        .annotate(d=TruncDate('uploaded_at'))
        .values('d')
        .annotate(c=Count('id'))
    )
    upload_counts = {row['d']: row['c'] for row in upload_rows}
    calendar_weeks = _build_calendar_weeks(cal_year, cal_month, upload_counts)
    calendar_nav = _calendar_nav_query(request, cal_year, cal_month, period, current_program)

    selected_day = None
    selected_day_docs = []
    day_param = request.GET.get('day', '').strip()
    if day_param:
        try:
            selected_day = date.fromisoformat(day_param)
        except ValueError:
            selected_day = None
    # The local date: uploads are bucketed by it, and between midnight and 8 a.m.
    # in Manila the UTC date is still yesterday.
    today = timezone.localdate(now)
    if selected_day is None and month_start <= today <= month_end:
        selected_day = today
    if selected_day:
        selected_day_docs = list(
            docs_qs.filter(uploaded_at__date=selected_day)
            .order_by('-uploaded_at')[:5]
        )

    # Each card opens the list its number counts, program filter included.
    program_param = f'&program={current_program.pk}' if current_program else ''
    duplicate_repo_url = reverse('documents:repository') + '?duplicate=review' + program_param
    week_repo_url = reverse('documents:repository') + '?uploaded=7d' + program_param

    # The dashboard summary export is for Administrators only.
    export_kind = request.GET.get('export')
    if export_kind in ('excel', 'csv') and is_admin(request.user):
        metrics = {
            'total_documents': total_documents,
            'docs_last7': docs_last7,
            'docs_last30': docs_last30,
            'duplicate_review_count': duplicate_review_count,
            'cluster_count': cluster_count,
            'ai_processed_count': ai_processed_count,
            'ai_processed_pct': ai_processed_pct,
        }
        if export_kind == 'excel' and openpyxl_available():
            return _excel_export(
                now, metrics, monthly_labels, monthly_counts,
                file_format_labels, file_format_counts, program_overview, current_program,
            )
        return _csv_export(
            now, total_documents, duplicate_review_count,
            monthly_labels, monthly_counts,
            file_format_labels, file_format_counts, program_overview,
            current_program,
        )

    local_now = timezone.localtime(now)
    hour = local_now.hour
    if hour < 12:
        greeting = 'Good morning'
    elif hour < 18:
        greeting = 'Good afternoon'
    else:
        greeting = 'Good evening'

    # The toolbar read "Last 6 months · Apr 2026 - Sep 2026" beside a control
    # whose own menu says "Monthly (6 months)". The span was the widest thing in
    # the row at 282px and half of it repeated the control next to it. The dates
    # are the part the control does not give, so the toolbar shows those and the
    # full phrase stays where it is still needed -- the chart's description for
    # a screen reader, which has no control beside it to borrow from.
    if period == 'daily':
        period_range = f"{trend_starts[0].strftime('%b %d')} – {now.strftime('%b %d')}"
        period_label = f"Last 30 days · {period_range}"
        period_display = 'Daily'
    elif period == 'weekly':
        period_range = f"{trend_starts[0].strftime('%b %d')} – {now.strftime('%b %d')}"
        period_label = f"Last 12 weeks · {period_range}"
        period_display = 'Weekly'
    else:
        period_range = f"{trend_starts[0].strftime('%b %Y')} – {now.strftime('%b %Y')}"
        period_label = f"Last 6 months · {period_range}"
        period_display = 'Monthly'

    context = {
        'total_documents': total_documents,
        'duplicate_review_count': duplicate_review_count,
        'duplicate_repo_url': duplicate_repo_url,
        'week_repo_url': week_repo_url,
        'recent_uploads': recent_uploads,
        'recent_activity': recent_activity,
        'docs_delta_pct': docs_delta['pct'],
        'docs_delta_dir': docs_delta['dir'],
        'docs_last30': docs_last30,
        'docs_last7': docs_last7,
        'week_delta_pct': week_delta['pct'],
        'week_delta_dir': week_delta['dir'],
        'cluster_count': cluster_count,
        # Descriptive name of each cluster number, shown on hover wherever the number is.
        'cluster_names': cluster_label_map(),
        'ai_processed_count': ai_processed_count,
        'ai_processed_pct': ai_processed_pct,
        'spark_counts': spark_counts,
        'monthly_labels': monthly_labels,
        'monthly_counts': monthly_counts,
        'monthly_total_window': monthly_total_window,
        'monthly_program_lines': program_lines,
        'file_format_labels': file_format_labels,
        'file_format_counts': file_format_counts,
        'file_format_total': file_format_total,
        'file_format_labels_count': len(file_format_labels),
        'period_label': period_label,
        'period_range': period_range,
        'period_display': period_display,
        'period_value': period,
        'all_programs': active_programs,
        'current_program': current_program,
        'calendar_year': cal_year,
        'calendar_month': cal_month,
        'calendar_month_label': date(cal_year, cal_month, 1).strftime('%B %Y'),
        'calendar_weeks': calendar_weeks,
        'calendar_nav': calendar_nav,
        'calendar_upload_counts_json': json.dumps(
            {d.isoformat(): c for d, c in upload_counts.items()}
        ),
        'selected_day': selected_day,
        'selected_day_docs': selected_day_docs,
        'selected_day_count': upload_counts.get(selected_day, 0) if selected_day else 0,
        'greeting': greeting,
        'generated_at': local_now,
    }
    return render(request, 'dashboard/dashboard.html', context)
