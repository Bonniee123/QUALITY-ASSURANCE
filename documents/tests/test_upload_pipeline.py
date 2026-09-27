"""
The upload path, after the performance work.

What these pin down, each measured before the change on a 32-file batch:

* The upload request extracts no text. It used to OCR every file before saving
  any of them -- 63 OCR calls and 99 seconds with the browser on "Finalizing on
  server…".
* Each saved file is extracted once, and its text is stored once. Images were
  OCR'd up to five times and kept identical text in two fields.
* Near-duplicate text is still refused. The check moved from before the save to
  after extraction; the verdict, the threshold and the upload-order semantics
  are unchanged, and a refused file is reported as a duplicate, not lost.
* Clustering runs once per batch and never overlaps itself, however many
  batches arrive; a batch reads as finished only once its run has finished.
* The embedding model is not asked to encode text it has already encoded.
"""
import importlib.util
import io
import threading
import time
from datetime import timedelta
from unittest import mock

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from documents import jobs
from documents.models import BackgroundJob, Document
from documents.upload_batches import STALL_AFTER, batch_detail
from documents.views import _is_near_duplicate_upload
from notifications.user_messages import bulk_upload_notification_message

BODY = ('The quality assurance office reviewed the evidence for area one and '
        'recorded its findings for the accreditation visit. ') * 4


def png_bytes(color=(40, 120, 200), quality_marker=0):
    """A small valid PNG; `quality_marker` changes the bytes but not the picture."""
    from PIL import Image
    img = Image.new('RGB', (64, 48), color)
    for x in range(0, 64, 8):
        for y in range(48):
            img.putpixel((x, y), (250, 250, 250))
    buf = io.BytesIO()
    img.save(buf, 'PNG', compress_level=1 + quality_marker)
    return buf.getvalue()


def docx_bytes(text):
    """A real (minimal) Word document: uploads are checked to be what they claim."""
    from docx import Document as WordDocument
    buf = io.BytesIO()
    word = WordDocument()
    word.add_paragraph(text)
    word.save(buf)
    return buf.getvalue()


def make_user(username='uploader', role='qa_staff'):
    user = User.objects.create_user(username, f'{username}@e.com', 'testpass123')
    user.profile.role = role
    user.profile.save()
    return user


def stored_doc(title, name, file_type='pdf', payload=b'%PDF-1.4 x', **extra):
    return Document.objects.create(
        title=title, file=SimpleUploadedFile(name, payload), file_type=file_type,
        year=2026, document_type='Report', **extra,
    )


class UploadRequestDoesNoExtractionTests(TestCase):
    """The request saves and hashes; reading the text is background work."""

    def setUp(self):
        self.user = make_user()
        self.client.login(username='uploader', password='testpass123')

    def post(self, files):
        return self.client.post(
            reverse('documents:bulk_upload'), {'files': files},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest', HTTP_ACCEPT='application/json',
        )

    def test_no_text_is_extracted_while_the_browser_waits(self):
        files = [
            SimpleUploadedFile('figure.png', png_bytes(), content_type='image/png'),
            SimpleUploadedFile('memo.pdf', b'%PDF-1.4 memo', content_type='application/pdf'),
            SimpleUploadedFile('minutes.docx', docx_bytes('Minutes of the meeting.'),
                               content_type='application/octet-stream'),
        ]
        refuse = AssertionError('text was extracted inside the upload request')
        with mock.patch('documents.jobs._dispatch_job_async'), \
                mock.patch('documents.views.extract_text', side_effect=refuse) as view_extract, \
                mock.patch('documents.text_extraction.extract_text_from_image', side_effect=refuse), \
                mock.patch('documents.text_extraction.extract_text_from_scanned_pdf', side_effect=refuse):
            response = self.post(files)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['uploaded_count'], 3)
        view_extract.assert_not_called()

    def test_each_file_is_saved_immediately_with_its_hashes(self):
        with mock.patch('documents.jobs._dispatch_job_async'):
            self.post([SimpleUploadedFile('figure.png', png_bytes(), content_type='image/png')])
        doc = Document.objects.get(file_type='png')
        self.assertFalse(doc.is_processed, 'saved, and waiting for background processing')
        self.assertTrue(doc.content_sha256)
        self.assertTrue(doc.image_phash, 'the visual hash the request computed is kept')

    def test_the_batch_job_lists_every_saved_file(self):
        with mock.patch('documents.jobs._dispatch_job_async'):
            response = self.post([
                SimpleUploadedFile('a.pdf', b'%PDF-1.4 a', content_type='application/pdf'),
                SimpleUploadedFile('b.pdf', b'%PDF-1.4 b', content_type='application/pdf'),
            ])
        job = BackgroundJob.objects.get(pk=response.json()['job_id'])
        self.assertEqual(len(job.payload['document_ids']), 2)

    def test_an_exact_copy_is_still_refused_at_upload(self):
        stored_doc('Original', 'original.pdf', payload=b'%PDF-1.4 same bytes')
        with mock.patch('documents.jobs._dispatch_job_async'):
            response = self.post([SimpleUploadedFile('copy.pdf', b'%PDF-1.4 same bytes',
                                                     content_type='application/pdf')])
        self.assertEqual(response.status_code, 400)
        self.assertIn('copy.pdf', response.json()['rejected_filenames'])

    def test_two_look_alike_images_in_one_batch_are_both_accepted_as_before(self):
        """
        Storing the visual hash at save time must not quietly start refusing the
        second of two look-alike images in the same batch: those have always
        been accepted and then flagged for review by the pipeline.
        """
        with mock.patch('documents.jobs._dispatch_job_async'):
            response = self.post([
                SimpleUploadedFile('fig.png', png_bytes(quality_marker=0), content_type='image/png'),
                SimpleUploadedFile('fig_1.png', png_bytes(quality_marker=5), content_type='image/png'),
            ])
        self.assertEqual(response.json()['uploaded_count'], 2)


class BackgroundProcessingTests(TestCase):
    """What the batch job does once the files are saved."""

    def setUp(self):
        self.user = make_user()
        self.extract_calls = []

    def fake_extract(self, texts):
        def extract(path):
            self.extract_calls.append(path)
            for marker, result in texts.items():
                if marker in str(path):
                    return result
            return '', 'pdf_extraction'
        return extract

    def run_batch(self, docs, texts):
        job = BackgroundJob.objects.create(
            job_type='bulk_upload_process', created_by=self.user,
            payload={'document_ids': [d.pk for d in docs], 'total_count': len(docs)},
        )
        with mock.patch('documents.jobs.extract_text', side_effect=self.fake_extract(texts)), \
                mock.patch('documents.jobs.request_full_pipeline') as pipeline:
            jobs._run_bulk_upload_process(job)
        job.refresh_from_db()
        return job, pipeline

    def test_each_file_is_extracted_exactly_once(self):
        docs = [stored_doc(f'upload {i}', f'upload-{i}.pdf') for i in range(5)]
        # Short, distinct text: identical long text would (rightly) make four of
        # the five near-duplicates of the first.
        self.run_batch(docs, {f'upload-{i}': (f'Area memo {i}', 'pdf_extraction') for i in range(5)})
        self.assertEqual(len(self.extract_calls), 5)
        self.assertEqual(Document.objects.filter(is_processed=True).count(), 5)

    def test_ocr_text_is_stored_once_not_in_both_fields(self):
        doc = stored_doc('scan', 'scan.png', file_type='png', payload=png_bytes())
        self.run_batch([doc], {'scan': ('Surf. Explore. Protect.', 'ocr_extraction')})
        doc.refresh_from_db()
        self.assertEqual(doc.ocr_text, 'Surf. Explore. Protect.')
        self.assertEqual(doc.extracted_text, '', 'the same text used to be stored twice')

    def test_text_already_extracted_is_reused(self):
        doc = stored_doc('already read', 'already.pdf', extracted_text=BODY)
        self.run_batch([doc], {'already': ('SHOULD NOT BE READ', 'pdf_extraction')})
        self.assertEqual(self.extract_calls, [])

    def test_clustering_is_requested_once_for_the_whole_batch(self):
        docs = [stored_doc(f'u{i}', f'u{i}.pdf') for i in range(4)]
        job, pipeline = self.run_batch(docs, {f'u{i}': (f'Area memo {i}', 'pdf_extraction')
                                              for i in range(4)})
        pipeline.assert_called_once()
        self.assertIn('analysis_requested_at', job.payload)

    def test_a_failed_extraction_marks_only_that_file(self):
        good, bad = stored_doc('good', 'good.pdf'), stored_doc('bad', 'bad.pdf')

        def extract(path):
            if 'bad' in str(path):
                raise OSError('unreadable file')
            return BODY, 'pdf_extraction'
        job = BackgroundJob.objects.create(
            job_type='bulk_upload_process', created_by=self.user,
            payload={'document_ids': [good.pk, bad.pk], 'total_count': 2})
        with mock.patch('documents.jobs.extract_text', side_effect=extract), \
                mock.patch('documents.jobs.request_full_pipeline'):
            jobs._run_bulk_upload_process(job)
        good.refresh_from_db()
        bad.refresh_from_db()
        self.assertTrue(good.is_processed)
        self.assertIn('unreadable', bad.processing_error)


class NearDuplicateAfterExtractionTests(TestCase):
    """The pre-save text check, now made after extraction -- same verdicts."""

    def setUp(self):
        self.user = make_user()
        self.original = stored_doc('Original Chapter', 'original.pdf',
                                   extracted_text=BODY, is_processed=True)

    def run_batch(self, docs, text_for):
        job = BackgroundJob.objects.create(
            job_type='bulk_upload_process', created_by=self.user,
            payload={'document_ids': [d.pk for d in docs], 'total_count': len(docs)},
        )
        with mock.patch('documents.jobs.extract_text',
                        side_effect=lambda path: (text_for(path), 'pdf_extraction')), \
                mock.patch('documents.jobs.request_full_pipeline'):
            jobs._run_bulk_upload_process(job)
        job.refresh_from_db()
        return job

    def test_a_near_duplicate_is_not_added(self):
        copy = stored_doc('copy', 'copy.pdf')
        name = copy.file.name
        storage = copy.file.storage
        job = self.run_batch([copy], lambda path: BODY)
        self.assertFalse(Document.objects.filter(pk=copy.pk).exists())
        self.assertFalse(storage.exists(name), 'its file is removed as well')
        self.assertTrue(Document.objects.filter(pk=self.original.pk).exists())
        self.assertIn(str(copy.pk), job.payload['rejected'])

    def test_the_upload_page_reports_it_as_a_duplicate_with_the_reason(self):
        copy = stored_doc('copy', 'copy.pdf')
        job = self.run_batch([copy], lambda path: BODY)
        row = batch_detail(job)['files'][0]
        self.assertEqual(row['state'], 'duplicate')
        self.assertIn('Original Chapter', row['detail'])
        self.assertEqual(row['name'], 'copy.pdf')

    def test_the_notification_does_not_count_it_as_uploaded(self):
        kept = stored_doc('kept', 'kept.pdf')
        copy = stored_doc('copy', 'copy.pdf')
        other = ('An unrelated laboratory safety inspection report with its own '
                 'findings and recommendations for the science building. ') * 4
        job = self.run_batch([kept, copy],
                             lambda path: other if 'kept' in str(path) else BODY)
        msg = bulk_upload_notification_message(job, success=True)
        self.assertIn('Your document was uploaded successfully', msg)
        self.assertIn('1 file matches an existing document', msg)

    def test_within_a_batch_the_later_file_is_the_one_refused(self):
        first = stored_doc('first', 'first.pdf')
        second = stored_doc('second', 'second.pdf')
        fresh = ('A new accreditation narrative that is not in the archive yet, '
                 'describing the programme outcomes and their assessment. ') * 4
        self.run_batch([first, second], lambda path: fresh)
        self.assertTrue(Document.objects.filter(pk=first.pk).exists())
        self.assertFalse(Document.objects.filter(pk=second.pk).exists())

    def test_a_different_document_is_kept(self):
        other = stored_doc('other', 'other.pdf')
        text = ('An unrelated laboratory safety inspection report with its own '
                'findings and recommendations for the science building. ') * 4
        self.run_batch([other], lambda path: text)
        other.refresh_from_db()
        self.assertTrue(other.is_processed)

    def test_short_text_is_never_judged_on_text(self):
        """Below the minimum length the text check abstains, as it always has."""
        tiny = stored_doc('tiny', 'tiny.pdf')
        self.run_batch([tiny], lambda path: 'Area I')
        self.assertTrue(Document.objects.filter(pk=tiny.pk).exists())

    def _faculty_in(self, area_code):
        from qa_structure.models import AccreditationArea
        own, _ = AccreditationArea.objects.get_or_create(area_code=area_code, defaults={'area_name': area_code})
        faculty = make_user(f'fac_{area_code.replace(" ", "")}', role='faculty')
        faculty.profile.assigned_areas.add(own)
        return faculty

    def test_faculty_are_not_shown_the_title_of_another_areas_document(self):
        from qa_structure.models import AccreditationArea
        other_area, _ = AccreditationArea.objects.get_or_create(area_code='Area II', defaults={'area_name': 'Faculty'})
        Document.objects.filter(pk=self.original.pk).update(acc_area=other_area, qa_area='Area II')
        self.user = self._faculty_in('Area III')
        copy = stored_doc('copy', 'copy.pdf')
        job = self.run_batch([copy], lambda path: BODY)
        reason = job.payload['rejected'][str(copy.pk)]['reason']
        self.assertNotIn('Original Chapter', reason)
        self.assertIn('a document already in the archive', reason)
        self.assertFalse(Document.objects.filter(pk=copy.pk).exists(), 'it is still refused')

    def test_faculty_see_the_title_when_the_match_is_in_their_own_area(self):
        from qa_structure.models import AccreditationArea
        area, _ = AccreditationArea.objects.get_or_create(area_code='Area III', defaults={'area_name': 'Curriculum'})
        Document.objects.filter(pk=self.original.pk).update(acc_area=area, qa_area='Area III')
        self.user = self._faculty_in('Area III')
        copy = stored_doc('copy', 'copy.pdf')
        job = self.run_batch([copy], lambda path: BODY)
        self.assertIn('Original Chapter', job.payload['rejected'][str(copy.pk)]['reason'])


class PreSaveCheckDoesNotReReadProcessedFilesTests(TestCase):
    """
    The single-file pages still check before saving. What stops is the check
    re-reading stored files that had already been processed and simply hold
    little text -- a photo was OCR'd again for every same-type upload.
    """

    def test_a_processed_low_text_image_is_not_read_again(self):
        stored_doc('booth', 'booth.png', file_type='png', payload=png_bytes(),
                   ocr_text='booth', is_processed=True)
        upload = SimpleUploadedFile('new.png', png_bytes(color=(9, 9, 9)))
        with mock.patch('documents.views.extract_text', return_value=(BODY, 'ocr_extraction')) as ext:
            _is_near_duplicate_upload(upload, '.png')
        self.assertEqual(ext.call_count, 1, 'only the upload itself is read')

    def test_a_backfilled_ocr_text_lands_in_the_ocr_field(self):
        doc = stored_doc('awaiting', 'awaiting.png', file_type='png', payload=png_bytes())
        upload = SimpleUploadedFile('new.png', png_bytes(color=(9, 9, 9)))
        with mock.patch('documents.views.extract_text', return_value=(BODY, 'ocr_extraction')):
            _is_near_duplicate_upload(upload, '.png')
        doc.refresh_from_db()
        self.assertEqual(doc.ocr_text, BODY.strip())
        self.assertEqual(doc.extracted_text, '')


class BatchFinishesWithItsAnalysisTests(TestCase):
    """Files turn Completed one by one; the batch waits for its clustering run."""

    def setUp(self):
        self.user = make_user()
        self.doc = stored_doc('done', 'done.pdf', extracted_text=BODY, is_processed=True)

    def job(self, requested_at):
        return BackgroundJob.objects.create(
            job_type='bulk_upload_process', created_by=self.user, status='completed',
            payload={'document_ids': [self.doc.pk], 'total_count': 1,
                     'analysis_requested_at': requested_at.isoformat()},
        )

    def test_not_finished_while_its_clustering_run_is_pending(self):
        batch = batch_detail(self.job(timezone.now()))
        self.assertEqual(batch['files'][0]['state'], 'completed')
        self.assertTrue(batch['analysing'])
        self.assertFalse(batch['finished'])

    def test_finished_once_a_run_started_after_the_request_completes(self):
        requested = timezone.now()
        job = self.job(requested)
        BackgroundJob.objects.create(job_type='full_ai_pipeline', status='completed',
                                     started_at=requested + timedelta(seconds=1))
        self.assertTrue(batch_detail(job)['finished'])

    def test_a_run_that_started_earlier_does_not_count(self):
        requested = timezone.now()
        job = self.job(requested)
        BackgroundJob.objects.create(job_type='full_ai_pipeline', status='completed',
                                     started_at=requested - timedelta(seconds=5))
        self.assertFalse(batch_detail(job)['finished'])

    def test_a_run_that_never_comes_is_reported_stalled(self):
        batch = batch_detail(self.job(timezone.now() - STALL_AFTER - timedelta(minutes=1)))
        self.assertTrue(batch['stalled'])


class PipelineCoalescingTests(SimpleTestCase):
    """At most one run at a time; everything asked for meanwhile shares one more."""

    def setUp(self):
        jobs._pipeline_running = False
        jobs._pipeline_dirty = False

    def tearDown(self):
        self.wait_idle()
        jobs._pipeline_running = False
        jobs._pipeline_dirty = False

    def wait_idle(self, timeout=5.0):
        deadline = time.monotonic() + timeout
        while jobs._pipeline_running and time.monotonic() < deadline:
            time.sleep(0.01)
        self.assertFalse(jobs._pipeline_running, 'the worker never went idle')

    def patched(self, fake_run):
        fake_jobs = mock.patch.object(jobs, 'BackgroundJob')
        run = mock.patch.object(jobs, 'run_job', side_effect=fake_run)
        return fake_jobs, run

    def test_requests_during_a_run_fold_into_one_follow_up(self):
        started, release, calls = threading.Event(), threading.Event(), []

        def fake_run(job):
            calls.append(job)
            if len(calls) == 1:
                started.set()
                release.wait(5)

        fake_jobs, run = self.patched(fake_run)
        with fake_jobs, run:
            self.assertTrue(jobs.request_full_pipeline())
            self.assertTrue(started.wait(5))
            folded = [jobs.request_full_pipeline() for _ in range(5)]
            release.set()
            self.wait_idle()
        self.assertEqual(folded, [False] * 5)
        self.assertEqual(len(calls), 2, 'one run, plus exactly one follow-up')

    def test_a_single_request_runs_once(self):
        calls = []
        fake_jobs, run = self.patched(calls.append)
        with fake_jobs, run:
            jobs.request_full_pipeline()
            self.wait_idle()
        self.assertEqual(len(calls), 1)

    def test_a_request_after_the_worker_is_idle_starts_a_new_run(self):
        calls = []
        fake_jobs, run = self.patched(calls.append)
        with fake_jobs, run:
            jobs.request_full_pipeline()
            self.wait_idle()
            self.assertTrue(jobs.request_full_pipeline())
            self.wait_idle()
        self.assertEqual(len(calls), 2)


class _FakeModel:
    def __init__(self):
        self.encoded = []

    def encode(self, texts, **kwargs):
        import numpy as np
        self.encoded.append(list(texts))
        return np.array([[float(len(t)), float(sum(map(ord, t)) % 997), 1.0] for t in texts])


# The free-memory check is off here: these tests are about the vector cache and
# use a fake model, so how much memory the PC has at the time must not matter.
@override_settings(AI_CLUSTER_EMBEDDING_MODEL='fake-model', AI_CLUSTER_EMBEDDING_MIN_FREE_MB=0)
class EmbeddingReuseTests(SimpleTestCase):
    """Only text the model has not seen is sent to it."""

    def setUp(self):
        if importlib.util.find_spec('sentence_transformers') is None:
            self.skipTest('sentence-transformers not installed')
        from ai_processing import embedding_service
        self.svc = embedding_service
        self.saved = dict(self.svc._model_cache), self.svc._vector_cache.copy()
        self.svc._model_cache.clear()
        self.svc._vector_cache.clear()
        self.model = _FakeModel()

    def tearDown(self):
        self.svc._model_cache.clear()
        self.svc._model_cache.update(self.saved[0])
        self.svc._vector_cache.clear()
        self.svc._vector_cache.update(self.saved[1])

    def doc(self, title):
        return Document(title=title, document_type='Report', year=2026, extracted_text=f'{title} body')

    def test_unchanged_documents_are_not_encoded_again(self):
        a, b, c = self.doc('alpha'), self.doc('beta'), self.doc('gamma')
        with mock.patch.object(self.svc, '_load_model', return_value=self.model):
            first = self.svc.compute_embedding_matrix([a, b])
            second = self.svc.compute_embedding_matrix([a, b, c])
        self.assertEqual(len(self.model.encoded[0]), 2)
        self.assertEqual(len(self.model.encoded[1]), 1, 'only the new document is encoded')
        self.assertTrue((first == second[:2]).all(), 'reused vectors are the same vectors')

    def test_changed_text_is_encoded_afresh(self):
        a = self.doc('alpha')
        with mock.patch.object(self.svc, '_load_model', return_value=self.model):
            self.svc.compute_embedding_matrix([a])
            a.extracted_text = 'alpha body, revised'
            self.svc.compute_embedding_matrix([a])
        self.assertEqual(len(self.model.encoded), 2)

    def test_rows_come_back_in_document_order(self):
        a, b = self.doc('alpha'), self.doc('beta')
        with mock.patch.object(self.svc, '_load_model', return_value=self.model):
            ab = self.svc.compute_embedding_matrix([a, b])
            ba = self.svc.compute_embedding_matrix([b, a])
        self.assertTrue((ab[0] == ba[1]).all() and (ab[1] == ba[0]).all())

    def test_the_model_is_loaded_from_local_files_first(self):
        from sentence_transformers import SentenceTransformer  # noqa: F401
        with mock.patch('sentence_transformers.SentenceTransformer',
                        return_value=self.model) as ctor:
            self.svc._load_model('fake-model')
        ctor.assert_called_once_with('fake-model', local_files_only=True)


class InterruptedJobsTests(TestCase):
    """
    A server restart must not leave a batch "running" for ever.

    Jobs run on threads inside the server process; a restart killed them and the
    row stayed 'running': its files sat in "Processing..." indefinitely (a
    12-file batch stopped at 0/12 never moved again).
    """

    def test_a_job_left_running_is_run_again_at_startup(self):
        user = make_user()
        doc = stored_doc('interrupted', 'interrupted.pdf', uploaded_by=user)
        job = BackgroundJob.objects.create(
            job_type='bulk_upload_process', status='running', created_by=user, started_at=timezone.now(),
            payload={'document_ids': [doc.pk], 'filenames': ['interrupted.pdf'], 'total_count': 1,
                     'processed_count': 0, 'errors': []})
        with mock.patch.object(jobs, '_extract_for_upload', return_value=(BODY, 'pypdf2', None)), \
                mock.patch.object(jobs, 'request_full_pipeline'):
            resumed = jobs.resume_pending_jobs()
        job.refresh_from_db()
        doc.refresh_from_db()
        self.assertEqual(resumed, 1)
        self.assertEqual(job.status, 'completed')
        self.assertEqual(job.payload['processed_count'], 1)
        self.assertTrue(doc.is_processed)

    def test_the_server_resumes_jobs_with_and_without_the_autoreloader(self):
        """--noreload has a single process and no RUN_MAIN; it never resumed jobs."""
        import os
        from qa_archiving_system.management.commands import runserver

        cases = (
            (True, None, False),     # autoreloader's file-watching parent: leave it to the child
            (True, 'true', True),    # autoreloader's serving child
            (False, None, True),     # --noreload: the only process
        )
        for use_reloader, run_main, starts in cases:
            env = dict(os.environ)
            env.pop('RUN_MAIN', None)
            if run_main:
                env['RUN_MAIN'] = run_main
            with mock.patch.dict(os.environ, env, clear=True), \
                    mock.patch.object(runserver.threading, 'Thread') as thread:
                runserver.start_pending_job_worker(use_reloader)
            self.assertEqual(thread.called, starts, (use_reloader, run_main))

    def test_finished_jobs_are_left_alone(self):
        done = BackgroundJob.objects.create(job_type='full_ai_pipeline', status='completed', result='ok')
        failed = BackgroundJob.objects.create(job_type='full_ai_pipeline', status='failed', error='x')
        self.assertEqual(jobs.recover_interrupted_jobs(), 0)
        done.refresh_from_db()
        failed.refresh_from_db()
        self.assertEqual((done.status, failed.status), ('completed', 'failed'))


class MatchedDocumentLabelTests(TestCase):
    """Visual-match messages name the matched document only to someone who may open it."""

    def setUp(self):
        from qa_structure.models import AccreditationArea
        self.area2, _ = AccreditationArea.objects.get_or_create(area_code='Area II', defaults={'area_name': 'Faculty'})
        self.area3, _ = AccreditationArea.objects.get_or_create(area_code='Area III', defaults={'area_name': 'Curriculum'})
        self.doc = stored_doc('Faculty Development Plan', 'plan.pdf', acc_area=self.area2, qa_area='Area II')

    def test_staff_see_the_title(self):
        from documents.views import matched_document_label
        self.assertEqual(matched_document_label(make_user('head', 'qa_staff'), self.doc), '"Faculty Development Plan"')

    def test_faculty_of_another_area_do_not(self):
        from documents.views import matched_document_label
        faculty = make_user('fac3', 'faculty')
        faculty.profile.assigned_areas.add(self.area3)
        self.assertEqual(matched_document_label(faculty, self.doc), 'a document already in the archive')

    def test_faculty_of_the_same_area_do(self):
        from documents.views import matched_document_label
        faculty = make_user('fac2', 'faculty')
        faculty.profile.assigned_areas.add(self.area2)
        self.assertEqual(matched_document_label(faculty, self.doc), '"Faculty Development Plan"')
