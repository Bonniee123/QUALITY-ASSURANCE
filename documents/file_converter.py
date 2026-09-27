"""
Convert DOCX and XLSX files to HTML for in-browser viewing.
Preserves formatting: headings, bold, italic, tables, lists, etc.
"""
import os
import logging
from html import escape

logger = logging.getLogger(__name__)


def docx_to_html(file_path):
    """
    Convert a DOCX file to formatted HTML.
    Preserves: headings, bold, italic, underline, tables, lists, alignment.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return '<p>python-docx is required for DOCX viewing.</p>'

    try:
        doc = Document(file_path)
        html_parts = []
        list_open = None  # Track if a list is currently open

        for element in doc.element.body:
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag

            if tag == 'p':
                # It's a paragraph
                para = None
                for p in doc.paragraphs:
                    if p._element is element:
                        para = p
                        break
                if para is None:
                    continue

                # Check if it's a list item
                numPr = para._element.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
                is_list = numPr is not None

                if is_list:
                    if list_open is None:
                        html_parts.append('<ul class="doc-list">')
                        list_open = 'ul'
                    html_parts.append(f'<li>{_runs_to_html(para.runs)}</li>')
                    continue
                else:
                    if list_open:
                        html_parts.append(f'</{list_open}>')
                        list_open = None

                # Determine HTML tag from style
                style_name = para.style.name.lower() if para.style else ''
                text_html = _runs_to_html(para.runs)

                if not text_html.strip():
                    html_parts.append('<div class="doc-spacer"></div>')
                    continue

                # Alignment
                align = ''
                if para.alignment == WD_ALIGN_PARAGRAPH.CENTER:
                    align = ' style="text-align:center"'
                elif para.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                    align = ' style="text-align:right"'
                elif para.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                    align = ' style="text-align:justify"'

                if 'heading 1' in style_name:
                    html_parts.append(f'<h1 class="doc-h1"{align}>{text_html}</h1>')
                elif 'heading 2' in style_name:
                    html_parts.append(f'<h2 class="doc-h2"{align}>{text_html}</h2>')
                elif 'heading 3' in style_name:
                    html_parts.append(f'<h3 class="doc-h3"{align}>{text_html}</h3>')
                elif 'heading 4' in style_name:
                    html_parts.append(f'<h4 class="doc-h4"{align}>{text_html}</h4>')
                elif 'title' in style_name:
                    html_parts.append(f'<h1 class="doc-title"{align}>{text_html}</h1>')
                elif 'subtitle' in style_name:
                    html_parts.append(f'<h3 class="doc-subtitle"{align}>{text_html}</h3>')
                else:
                    html_parts.append(f'<p class="doc-para"{align}>{text_html}</p>')

            elif tag == 'tbl':
                # It's a table
                if list_open:
                    html_parts.append(f'</{list_open}>')
                    list_open = None

                for table in doc.tables:
                    if table._element is element:
                        html_parts.append(_table_to_html(table))
                        break

        # Close any open list
        if list_open:
            html_parts.append(f'</{list_open}>')

        return '\n'.join(html_parts)

    except Exception as e:
        logger.error(f"DOCX to HTML error: {e}")
        return f'<p class="doc-error">Error rendering document: {escape(str(e))}</p>'


def _runs_to_html(runs):
    """Convert paragraph runs to HTML with inline formatting."""
    parts = []
    for run in runs:
        text = escape(run.text)
        if not text:
            continue

        # Apply formatting
        if run.bold:
            text = f'<strong>{text}</strong>'
        if run.italic:
            text = f'<em>{text}</em>'
        if run.underline:
            text = f'<u>{text}</u>'
        if run.font.strike:
            text = f'<s>{text}</s>'

        # Font size
        style_parts = []
        if run.font.size:
            pt = run.font.size.pt
            style_parts.append(f'font-size:{pt}pt')

        # Font color
        if run.font.color and run.font.color.rgb:
            color = str(run.font.color.rgb)
            style_parts.append(f'color:#{color}')

        # Font name
        if run.font.name:
            style_parts.append(f"font-family:'{run.font.name}',sans-serif")

        if style_parts:
            text = f'<span style="{";".join(style_parts)}">{text}</span>'

        parts.append(text)

    return ''.join(parts)


def _table_to_html(table):
    """Convert a DOCX table to HTML."""
    rows_html = []
    for i, row in enumerate(table.rows):
        cells_html = []
        for cell in row.cells:
            cell_text = escape(cell.text.strip())
            if i == 0:
                cells_html.append(f'<th>{cell_text}</th>')
            else:
                cells_html.append(f'<td>{cell_text}</td>')
        tag = 'tr'
        rows_html.append(f'<tr>{"".join(cells_html)}</tr>')

    header = rows_html[0] if rows_html else ''
    body = '\n'.join(rows_html[1:]) if len(rows_html) > 1 else ''

    return f'<table class="doc-table"><thead>{header}</thead><tbody>{body}</tbody></table>'


def xlsx_to_html(file_path):
    """
    Convert an XLSX file to formatted HTML tables.
    Each sheet becomes a separate table with a header.
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        return '<p>openpyxl is required for XLSX viewing.</p>'

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        html_parts = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            html_parts.append(f'<h2 class="doc-h2 sheet-name"><i class="bi bi-table me-2"></i>{escape(sheet_name)}</h2>')
            html_parts.append('<div class="table-responsive"><table class="doc-table">')

            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = []
                for cell in row:
                    value = escape(str(cell)) if cell is not None else ''
                    if i == 0:
                        cells.append(f'<th>{value}</th>')
                    else:
                        cells.append(f'<td>{value}</td>')
                html_parts.append(f'<tr>{"".join(cells)}</tr>')

            html_parts.append('</table></div>')

        wb.close()
        return '\n'.join(html_parts)

    except Exception as e:
        logger.error(f"XLSX to HTML error: {e}")
        return f'<p class="doc-error">Error rendering spreadsheet: {escape(str(e))}</p>'
