"""Shared helpers for producing professionally styled Excel (.xlsx) exports.

Used by the reports, dashboard, and repository exporters so every spreadsheet
download in the system shares one consistent, branded look (colored title
band, colored header row, banded rows, borders, auto-sized columns, frozen
header, and an autofilter).
"""
from io import BytesIO

from django.http import HttpResponse

EXCEL_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Brand palette (hex, no leading '#').
_BRAND = '0E4D5C'        # dark teal title band
_HEADER = '1F7A8C'       # teal header row
_BAND = 'EEF4F6'         # light banded row
_BORDER = 'CBD8DE'
_TITLE_TEXT = 'FFFFFF'
_SUB_TEXT = '5A6B72'
_CELL_TEXT = '1F2A30'


# A cell starting with one of these is read by Excel as a formula. Titles, names
# and program codes are typed by users, so a document titled
# '=HYPERLINK("http://...","Click")' became a live link in every export.
_FORMULA_TRIGGERS = ('=', '+', '-', '@', '\t', '\r')


def csv_safe(value):
    """A CSV value Excel will show as text (a leading apostrophe on a formula-like string)."""
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGERS):
        return "'" + value
    return value


def csv_safe_row(row):
    return [csv_safe(value) for value in row]


def openpyxl_available():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def build_styled_sheet(ws, title, headers, rows, generated_at=None, subtitle=None):
    """Render a styled table onto an existing worksheet.

    Returns the 1-based index of the last written row.
    """
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    ws.sheet_view.showGridLines = False
    ncols = max(len(headers), 1)
    last_col = get_column_letter(ncols)

    brand_fill = PatternFill('solid', fgColor=_BRAND)
    header_fill = PatternFill('solid', fgColor=_HEADER)
    band_fill = PatternFill('solid', fgColor=_BAND)
    title_font = Font(name='Calibri', size=16, bold=True, color=_TITLE_TEXT)
    sub_font = Font(name='Calibri', size=10, italic=True, color=_SUB_TEXT)
    header_font = Font(name='Calibri', size=11, bold=True, color=_TITLE_TEXT)
    cell_font = Font(name='Calibri', size=10, color=_CELL_TEXT)
    thin = Side(style='thin', color=_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
    title_align = Alignment(horizontal='left', vertical='center', indent=1)

    r = 1
    ws.merge_cells(f'A{r}:{last_col}{r}')
    tc = ws.cell(row=r, column=1, value=title)
    tc.font = title_font
    tc.fill = brand_fill
    tc.alignment = title_align
    ws.row_dimensions[r].height = 30
    r += 1

    meta_lines = []
    if generated_at is not None:
        meta_lines.append(f'Generated: {generated_at:%Y-%m-%d %H:%M}')
        meta_lines.append(f'Total records: {len(rows)}')
    if subtitle:
        meta_lines.extend(subtitle if isinstance(subtitle, (list, tuple)) else [subtitle])
    for line in meta_lines:
        ws.merge_cells(f'A{r}:{last_col}{r}')
        mc = ws.cell(row=r, column=1, value=line)
        mc.font = sub_font
        mc.alignment = title_align
        r += 1

    r += 1  # spacer row
    header_row = r
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border
    ws.row_dimensions[r].height = 22
    r += 1

    data_start = r
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=ci, value=val)
            if isinstance(val, str) and val.startswith('='):
                # openpyxl stores a string starting with "=" as a formula; data
                # is written as text, whatever it starts with.
                c.data_type = 's'
            c.font = cell_font
            c.border = border
            c.alignment = left if ci == 1 else center
            if ri % 2 == 1:
                c.fill = band_fill
        r += 1

    last_row = r - 1
    ws.freeze_panes = f'A{data_start}'
    ws.auto_filter.ref = f'A{header_row}:{last_col}{max(header_row, last_row)}'

    for ci in range(1, ncols + 1):
        maxlen = len(str(headers[ci - 1]))
        for row in rows:
            val = row[ci - 1] if ci - 1 < len(row) else ''
            if val is not None:
                maxlen = max(maxlen, len(str(val)))
        ws.column_dimensions[get_column_letter(ci)].width = min(max(maxlen + 4, 12), 60)

    return last_row


def single_sheet_workbook(title, headers, rows, generated_at=None, sheet_name=None, subtitle=None):
    """Create a one-sheet styled workbook."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or title)[:31]
    build_styled_sheet(ws, title, headers, rows, generated_at=generated_at, subtitle=subtitle)
    return wb


def excel_response(wb, filename_base):
    """Serialize a workbook to an attachment HttpResponse."""
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=EXCEL_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename_base}.xlsx"'
    return response
