"""
Text extraction utilities for different file formats.
Supports PDF, DOCX, and XLSX. Image OCR is supported when Tesseract is installed.
"""
import os
import logging
import threading

from django.conf import settings

logger = logging.getLogger(__name__)


# MuPDF (PyMuPDF / ``fitz``) keeps process-wide state and is not safe to drive
# from several threads at once. Uploads are now extracted in parallel, so page
# rendering is serialised here. Rendering is the fast part; the OCR that follows
# runs in a separate Tesseract process and stays outside the lock.
_FITZ_LOCK = threading.Lock()


def _default_tesseract_cmd():
    """
    Return the Tesseract executable path:
    1) settings.TESSERACT_CMD if set, else
    2) env var TESSERACT_CMD, else
    3) the common Windows install path (only if it exists), else
    4) None — let pytesseract find it on PATH.
    """
    cmd = getattr(settings, 'TESSERACT_CMD', '') or os.getenv('TESSERACT_CMD', '')
    if cmd:
        return cmd
    win_default = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    if os.name == 'nt' and os.path.exists(win_default):
        return win_default
    return None


def extract_text_from_pdf(file_path):
    """
    Extract text from a machine-readable PDF file using pypdf.
    Returns extracted text as a string.
    """
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return '\n'.join(text_parts)
    except Exception as e:
        # Invalid or minimal PDFs (common in tests) surface as EOF / parse errors — avoid ERROR noise
        msg = str(e).lower()
        if 'eof' in msg or 'stream' in msg or 'parse' in msg:
            logger.warning('PDF extraction skipped (invalid or minimal PDF) for %s: %s', file_path, e)
        else:
            logger.error('PDF extraction error for %s: %s', file_path, e)
        return ''


def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX file using python-docx.
    Returns all paragraph text as a string.
    """
    try:
        from docx import Document
        doc = Document(file_path)
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text.strip())
        return '\n'.join(text_parts)
    except Exception as e:
        logger.error(f"DOCX extraction error for {file_path}: {e}")
        return ''


def extract_text_from_xlsx(file_path):
    """
    Extract text from an XLSX file using openpyxl.
    Returns all non-empty cell values as a string.
    """
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True)
        text_parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                row_text = ' '.join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text_parts.append(row_text.strip())
        wb.close()
        return '\n'.join(text_parts)
    except Exception as e:
        logger.error(f"XLSX extraction error for {file_path}: {e}")
        return ''


def _upright(img):
    """The image rotated as its EXIF orientation tag says (unchanged if it has none)."""
    try:
        from PIL import ImageOps
        return ImageOps.exif_transpose(img)
    except Exception:
        return img


def extract_text_from_image(file_path):
    """
    Extract text from an image file using Tesseract OCR.
    The Tesseract binary path can be configured via settings.TESSERACT_CMD or env TESSERACT_CMD.
    """
    try:
        import pytesseract
        from PIL import Image

        cmd = _default_tesseract_cmd()
        if cmd:
            pytesseract.pytesseract.tesseract_cmd = cmd

        with Image.open(file_path) as img:
            # Same handling as the PDF OCR path: the configured language, and the
            # picture turned upright first -- a phone photo stored sideways with
            # an EXIF orientation tag used to OCR as gibberish.
            text = pytesseract.image_to_string(_upright(img), lang=getattr(settings, 'OCR_LANGUAGE', 'eng'))

        return text.strip() if text and text.strip() else ''
    except ImportError:
        logger.warning("pytesseract or Pillow not installed. OCR disabled.")
        return ''
    except Exception as e:
        logger.error(f"OCR extraction error for {file_path}: {e}")
        return ''


def extract_text_from_scanned_pdf(file_path, max_pages=6):
    """
    OCR fallback for scanned/image-only PDFs.
    Tries PyMuPDF first, then pdf2image if available.
    """
    try:
        import pytesseract
    except ImportError:
        logger.warning('pytesseract not installed. PDF OCR fallback disabled.')
        return ''

    cmd = _default_tesseract_cmd()
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd

    texts = []
    page_limit = max(1, int(max_pages or 1))
    ocr_lang = getattr(settings, 'OCR_LANGUAGE', 'eng')

    # Path 1: PyMuPDF
    try:
        import fitz  # type: ignore
        import io
        from PIL import Image

        with _FITZ_LOCK:
            with fitz.open(file_path) as doc:
                pages = [
                    doc.load_page(page_index).get_pixmap(dpi=200, alpha=False).tobytes('png')
                    for page_index in range(min(page_limit, len(doc)))
                ]
        for png in pages:
            img = Image.open(io.BytesIO(png))
            text = pytesseract.image_to_string(img, lang=ocr_lang).strip()
            if text:
                texts.append(text)
        if texts:
            return '\n'.join(texts)
    except Exception as e:
        logger.info('PyMuPDF OCR fallback unavailable for %s: %s', file_path, e)

    # Path 2: pdf2image
    try:
        from pdf2image import convert_from_path

        poppler_path = (getattr(settings, 'POPPLER_PATH', '') or os.getenv('POPPLER_PATH', '')).strip() or None
        images = convert_from_path(
            file_path,
            dpi=200,
            first_page=1,
            last_page=page_limit,
            poppler_path=poppler_path,
        )
        for img in images:
            text = pytesseract.image_to_string(img, lang=ocr_lang).strip()
            if text:
                texts.append(text)
        if texts:
            return '\n'.join(texts)
    except Exception as e:
        logger.info('pdf2image OCR fallback unavailable for %s: %s', file_path, e)

    return ''


def extract_text(file_path):
    """
    Auto-detect file type and extract text accordingly.
    Returns tuple: (extracted_text, method_used)
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        text = extract_text_from_pdf(file_path)
        if text and len(text.strip()) >= int(getattr(settings, 'PDF_OCR_MIN_TEXT_CHARS', 120)):
            return text, 'pdf_extraction'

        enable_pdf_ocr = bool(getattr(settings, 'ENABLE_PDF_OCR_FALLBACK', True))
        max_pdf_ocr_mb = int(getattr(settings, 'PDF_OCR_MAX_FILE_MB', 20))
        max_pdf_ocr_pages = int(getattr(settings, 'PDF_OCR_MAX_PAGES', 6))
        try:
            file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
        except OSError:
            file_size_mb = 0

        if enable_pdf_ocr and file_size_mb <= max_pdf_ocr_mb:
            ocr_text = extract_text_from_scanned_pdf(file_path, max_pages=max_pdf_ocr_pages)
            if ocr_text:
                return ocr_text, 'pdf_ocr_fallback'
        return text, 'pdf_extraction'
    elif ext == '.docx':
        text = extract_text_from_docx(file_path)
        return text, 'docx_extraction'
    elif ext == '.xlsx':
        text = extract_text_from_xlsx(file_path)
        return text, 'xlsx_extraction'
    elif ext in ['.jpg', '.jpeg', '.png']:
        text = extract_text_from_image(file_path)
        return text, 'ocr_extraction' if text else 'ocr_failed'
    else:
        return '', 'unsupported_format'


def extract_text_detailed(file_path):
    """
    ``extract_text`` plus two plain-language explanations for the document page:

    * ``problem`` -- why no text came out, when that is a fault rather than a
      file with no words in it: a damaged Word or Excel file, an image too large
      to open, or OCR that cannot run because Tesseract is missing. Each of these
      used to end as a generic "no text" -- or, for a damaged Word file, as a
      success with nothing in it.
    * ``notice`` -- what was left out of a scanned PDF: OCR reads only the first
      PDF_OCR_MAX_PAGES pages and skips files over PDF_OCR_MAX_FILE_MB, and the
      rest was silently unsearchable.

    Returns ``(text, method, problem, notice)``.
    """
    text, method = extract_text(file_path)
    return (text, method) + explain_extraction(file_path, text, method)


def explain_extraction(file_path, text, method):
    """``(problem, notice)`` for a result ``extract_text`` already returned (see above)."""
    ext = os.path.splitext(file_path)[1].lower()
    problem = '' if (text or '').strip() else _why_no_text(file_path, ext)
    notice = _scanned_pdf_notice(file_path, method, text) if ext == '.pdf' else ''
    return problem, notice


def is_damaged_file_problem(problem):
    """True for a problem that means the file itself cannot be read (see _why_no_text)."""
    return bool(problem) and problem.startswith('Could not')


def ocr_unavailable_reason():
    """Why OCR cannot run on this server, or '' when Tesseract answers."""
    try:
        import pytesseract
    except ImportError:
        return 'OCR is not available: the pytesseract package is not installed on the server.'
    cmd = _default_tesseract_cmd()
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        return ('OCR is not available: the Tesseract program was not found on the server '
                '(check TESSERACT_CMD), so no text could be read from this file.')
    return ''


def _why_no_text(file_path, ext):
    """The fault behind an empty extraction, or '' when the file simply has no text."""
    if ext == '.docx':
        try:
            from docx import Document
            Document(file_path)
        except Exception:
            return 'Could not read this Word document - the file may be damaged or incomplete.'
        return ''
    if ext == '.xlsx':
        try:
            from openpyxl import load_workbook
            load_workbook(file_path, read_only=True).close()
        except Exception:
            return 'Could not read this Excel workbook - the file may be damaged or incomplete.'
        return ''
    if ext in ('.jpg', '.jpeg', '.png'):
        try:
            from PIL import Image
            with Image.open(file_path) as img:
                img.load()
        except Exception as exc:
            if type(exc).__name__ == 'DecompressionBombError':
                return ('This image is too large to read (over Pillow\'s safety limit on pixels); '
                        'save it at a smaller size and upload it again.')
            return 'Could not open this image - the file may be damaged.'
        return ocr_unavailable_reason()
    if ext == '.pdf':
        if not _pdf_opens(file_path):
            return 'Could not read this PDF - the file may be damaged or incomplete.'
        if bool(getattr(settings, 'ENABLE_PDF_OCR_FALLBACK', True)) and not _pdf_too_big_for_ocr(file_path):
            return ocr_unavailable_reason()
    return ''


def _pdf_opens(file_path):
    """True if either PDF reader can open the file (pypdf is stricter than MuPDF)."""
    try:
        from pypdf import PdfReader
        PdfReader(file_path)
        return True
    except Exception:
        pass
    try:
        import fitz  # type: ignore
        with _FITZ_LOCK:
            with fitz.open(file_path) as doc:
                return doc.page_count > 0
    except Exception:
        return False


def _pdf_too_big_for_ocr(file_path):
    try:
        return os.path.getsize(file_path) / (1024 * 1024) > int(getattr(settings, 'PDF_OCR_MAX_FILE_MB', 20))
    except OSError:
        return False


def _pdf_page_count(file_path):
    try:
        from pypdf import PdfReader
        return len(PdfReader(file_path).pages)
    except Exception:
        return 0


def _scanned_pdf_notice(file_path, method, text):
    """What OCR left out of a scanned PDF, in words the uploader can act on."""
    max_mb = int(getattr(settings, 'PDF_OCR_MAX_FILE_MB', 20))
    max_pages = int(getattr(settings, 'PDF_OCR_MAX_PAGES', 6))
    if method == 'pdf_ocr_fallback':
        pages = _pdf_page_count(file_path)
        if pages > max_pages:
            return (f'Scanned PDF: only the first {max_pages} of its {pages} pages were read by OCR, '
                    f'so text on pages {max_pages + 1}-{pages} is not searchable.')
        return ''
    # extract_text returns 'pdf_extraction' with too little text exactly when the
    # text layer was too thin and OCR did not run or found nothing; the size is
    # what tells "skipped" from "found nothing".
    if (method == 'pdf_extraction' and bool(getattr(settings, 'ENABLE_PDF_OCR_FALLBACK', True))
            and len((text or '').strip()) < int(getattr(settings, 'PDF_OCR_MIN_TEXT_CHARS', 120))
            and _pdf_too_big_for_ocr(file_path)):
        return (f'Scanned PDF over {max_mb} MB: OCR was skipped, so its text is not searchable. '
                f'Split it into smaller files to make it searchable.')
    return ''
