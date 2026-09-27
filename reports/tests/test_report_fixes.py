"""
Reports count what the rest of the system counts (S-DR-04) and export text as
text (S-SEC-02).

Before: the Inventory and Cluster reports included archived versions, so their
totals disagreed with the Dashboard; the Recent report showed 50 rows and
exported 200; a document titled '=HYPERLINK(...)' became a live formula in the
CSV and Excel exports.
"""
import csv
import io

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from documents.models import ClusterResult, Document

FORMULA = '=HYPERLINK("http://evil.example","Click me")'


class ReportTests(TestCase):

    def setUp(self):
        self.admin = User.objects.create_user('rep_admin', password='pass12345')
        self.admin.profile.role = 'admin'
        self.admin.profile.save()
        self.client.force_login(self.admin)
        for i in range(55):
            Document.objects.create(title=f'Doc {i}', file=f'uploaded_documents/d{i}.pdf', file_type='pdf', year=2026,
                                    document_type='Report', cluster_label=1)
        self.archived = Document.objects.create(title='Archived version', file='uploaded_documents/old.pdf',
                                                file_type='pdf', year=2026, document_type='Report',
                                                cluster_label=1, is_archived=True)
        self.formula = Document.objects.create(title=FORMULA, file='uploaded_documents/f.pdf', file_type='pdf',
                                               year=2026, document_type='Report')
        ClusterResult.objects.create(document=Document.objects.first(), cluster_number=1,
                                     cluster_label='Area III · Report · minutes')

    def csv_rows(self, report_type):
        text = self.client.get(reverse('reports:export_csv') + f'?type={report_type}').content.decode()
        return list(csv.reader(io.StringIO(text)))

    def test_archived_versions_are_left_out(self):
        page = self.client.get(reverse('reports:page') + '?type=inventory')
        self.assertEqual(page.context['record_count'], 56)
        self.assertNotContains(page, 'Archived version')
        self.assertNotIn('Archived version', str(self.csv_rows('inventory')))

    def test_the_recent_export_has_the_rows_the_page_shows(self):
        page = self.client.get(reverse('reports:page') + '?type=recent')
        self.assertEqual(page.context['record_count'], 50)
        rows = self.csv_rows('recent')
        self.assertIn(['Total records: 50'], rows)

    def test_the_cluster_export_names_each_cluster_and_counts_current_documents(self):
        rows = self.csv_rows('cluster')
        self.assertIn(['Cluster', 'Cluster Name', 'Document Count'], rows)
        self.assertIn(['Cluster 1', 'Area III · Report · minutes', '55'], rows)

    def test_formula_titles_are_exported_as_text(self):
        titles = [row[0] for row in self.csv_rows('inventory') if row]
        self.assertIn("'" + FORMULA, titles)
        self.assertNotIn(FORMULA, titles)

    def test_the_excel_export_stores_the_title_as_text(self):
        from openpyxl import load_workbook
        data = self.client.get(reverse('reports:export_excel') + '?type=inventory').content
        sheet = load_workbook(io.BytesIO(data)).active
        cells = [c for row in sheet.iter_rows() for c in row if c.value == FORMULA]
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0].data_type, 's', 'a string, not a formula')
